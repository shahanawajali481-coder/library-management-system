"""
Library Management System (LMS) - Web Application
Author: SHAHANAWAJ (Roll No: 2410302051)
College Project | Flask + SQLite + Vanilla CSS & JS
"""

import os
import sys
import io
import csv
import sqlite3
import tempfile
import shutil
from datetime import datetime
from flask import (
    Flask,
    render_template,
    request,
    redirect,
    url_for,
    flash,
    jsonify,
    Response,
    session,
    send_file
)

from config import (
    APP_TITLE,
    APP_SUBTITLE,
    STUDENT_NAME,
    STUDENT_ROLL,
    CURRENCY_SYMBOL,
    FINE_PER_DAY,
    DEFAULT_LOAN_DAYS,
    DB_PATH
)
from database import init_db
from services.book_service import BookService
from services.member_service import MemberService
from services.transaction_service import TransactionService
from services.settings_service import SettingsService

# Initialize Flask application
app = Flask(__name__)
# Secure secret key for session management and flash messages
# Set SECRET_KEY in production; the fallback is only for local development.
app.secret_key = os.environ.get("SECRET_KEY") or os.urandom(32)
app.config["MAX_CONTENT_LENGTH"] = 50 * 1024 * 1024  # Limit uploaded database backups to 50 MB


# Context Processor to inject global variables into all Jinja templates
@app.context_processor
def inject_globals():
    settings = SettingsService.load_settings()
    # Check session theme preference, fallback to database setting
    current_theme = session.get("theme", "dark" if settings.get("dark_mode", True) else "light")
    current_user = None
    if session.get("user_id"):
        current_user = {
            "id": session.get("user_id"),
            "username": session.get("username", "admin"),
            "role": session.get("role", "Admin")
        }
    return {
        "app_title": APP_TITLE,
        "app_subtitle": APP_SUBTITLE,
        "student_name": STUDENT_NAME,
        "student_roll": STUDENT_ROLL,
        "currency_symbol": CURRENCY_SYMBOL,
        "fine_per_day": FINE_PER_DAY,
        "default_loan_days": DEFAULT_LOAN_DAYS,
        "current_theme": current_theme,
        "current_user": current_user,
        "now_year": datetime.now().year,
        "current_endpoint": request.endpoint or ""
    }


# Cache control to prevent browser back-button access after logout
@app.after_request
def add_cache_control_headers(response):
    """Ensure pages are not cached so browser Back button forces re-authentication after logout."""
    response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate, max-age=0"
    response.headers["Pragma"] = "no-cache"
    response.headers["Expires"] = "0"
    return response


# Authentication Guard
@app.before_request
def require_authentication():
    """Protect all dashboard and circulation routes from unauthenticated access."""
    open_endpoints = {"login", "signup", "static", "toggle_theme"}
    if request.endpoint and request.endpoint not in open_endpoints:
        if not session.get("user_id"):
            if request.endpoint != "index":
                flash("Please log in to access the library management system.", "warning")
            return redirect(url_for("login", next=request.path))


# =============================================================================
# 0. AUTHENTICATION & SESSION MANAGEMENT
# =============================================================================
@app.route("/signup", methods=["GET", "POST"])
def signup():
    """Sign up route for creating a new user account."""
    if session.get("user_id"):
        return redirect(url_for("dashboard"))

    if request.method == "POST":
        name = request.form.get("name", "").strip()
        username = request.form.get("username", "").strip()
        email = request.form.get("email", "").strip()
        password = request.form.get("password", "")
        confirm_password = request.form.get("confirm_password", "")

        if not name or not username or not email or not password or not confirm_password:
            flash("All fields are required. Please fill in the entire form.", "danger")
            return render_template("signup.html", form_data=request.form)

        if password != confirm_password:
            flash("Passwords do not match. Please ensure both password fields match.", "danger")
            return render_template("signup.html", form_data=request.form)

        success, msg, new_id = SettingsService.register_user(
            name=name,
            username=username,
            email=email,
            password=password,
            role="Staff"
        )

        if success:
            flash(msg, "success")
            return redirect(url_for("login"))
        else:
            flash(msg, "danger")
            return render_template("signup.html", form_data=request.form)

    return render_template("signup.html", form_data={})


@app.route("/login", methods=["GET", "POST"])
def login():
    if session.get("user_id"):
        return redirect(url_for("dashboard"))

    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        next_url = request.form.get("next") or request.args.get("next")

        user = SettingsService.verify_login(username, password)
        if user:
            session["user_id"] = user["id"]
            session["username"] = user["username"]
            session["role"] = user.get("role", "Admin")
            session["name"] = user.get("name") or user["username"]
            display_name = user.get("name") or user["username"]
            flash(f"Welcome back, {display_name}! Logged in successfully.", "success")
            if next_url and next_url.startswith("/") and not next_url.startswith("//"):
                return redirect(next_url)
            return redirect(url_for("dashboard"))
        else:
            flash("Invalid username/email or password. Please try again.", "danger")
            return render_template("login.html", username=username, next=next_url)

    next_url = request.args.get("next", "")
    return render_template("login.html", next=next_url)


@app.route("/logout", methods=["GET", "POST"])
def logout():
    """Clear session data and redirect to login with confirmation message."""
    session.pop("user_id", None)
    session.pop("username", None)
    session.pop("role", None)
    session.pop("name", None)
    flash("Logged out successfully.", "success")
    return redirect(url_for("login"))


# =============================================================================
# 1. DASHBOARD
# =============================================================================
@app.route("/")
def index():
    if session.get("user_id"):
        return redirect(url_for("dashboard"))
    return redirect(url_for("login"))


@app.route("/dashboard")
def dashboard():
    stats = TransactionService.get_dashboard_stats()
    return render_template("dashboard.html", stats=stats)


# =============================================================================
# 2. BOOKS MANAGEMENT
# =============================================================================
@app.route("/books")
def books_list():
    search_query = request.args.get("q", "").strip()
    category = request.args.get("category", "All")
    sort_by = request.args.get("sort", "title")

    books = BookService.get_all_books(search_query=search_query, category=category, sort_by=sort_by)
    categories = BookService.get_categories()

    # Calculate summary metrics
    total_titles = len(books)
    total_copies = sum(b["quantity"] for b in books)
    total_available = sum(b["available"] for b in books)

    return render_template(
        "books.html",
        books=books,
        categories=categories,
        search_query=search_query,
        selected_category=category,
        selected_sort=sort_by,
        total_titles=total_titles,
        total_copies=total_copies,
        total_available=total_available
    )


@app.route("/books/new", methods=["GET", "POST"])
def add_book():
    if request.method == "POST":
        isbn = request.form.get("isbn", "").strip()
        title = request.form.get("title", "").strip()
        author = request.form.get("author", "").strip()
        category = request.form.get("category", "").strip()
        quantity_str = request.form.get("quantity", "1").strip()

        try:
            quantity = int(quantity_str)
        except ValueError:
            flash("Quantity must be a valid positive integer.", "danger")
            categories = BookService.get_categories()
            return render_template("book_form.html", action="Add", book=request.form, categories=categories)

        success, msg, book_id = BookService.add_book(
            isbn=isbn,
            title=title,
            author=author,
            category=category,
            quantity=quantity
        )

        if success:
            flash(msg, "success")
            return redirect(url_for("books_list"))
        else:
            flash(msg, "danger")
            categories = BookService.get_categories()
            return render_template("book_form.html", action="Add", book=request.form, categories=categories)

    categories = BookService.get_categories()
    return render_template("book_form.html", action="Add", book={}, categories=categories)


@app.route("/books/<int:book_id>/edit", methods=["GET", "POST"])
def edit_book(book_id):
    book = BookService.get_book_by_id(book_id)
    if not book:
        flash("Book not found.", "danger")
        return redirect(url_for("books_list"))

    categories = BookService.get_categories()

    if request.method == "POST":
        isbn = request.form.get("isbn", "").strip()
        title = request.form.get("title", "").strip()
        author = request.form.get("author", "").strip()
        category = request.form.get("category", "").strip()
        quantity_str = request.form.get("quantity", "").strip()
        available_str = request.form.get("available", "").strip()

        try:
            quantity = int(quantity_str)
            available = int(available_str)
        except ValueError:
            flash("Quantity and Available copies must be valid non-negative integers.", "danger")
            return render_template("book_form.html", action="Edit", book=request.form, categories=categories, book_id=book_id)

        success, msg = BookService.update_book(
            book_id=book_id,
            isbn=isbn,
            title=title,
            author=author,
            category=category,
            quantity=quantity,
            available=available
        )

        if success:
            flash(msg, "success")
            return redirect(url_for("books_list"))
        else:
            flash(msg, "danger")
            # Preserve user input in the form
            form_data = dict(request.form)
            form_data["id"] = book_id
            return render_template("book_form.html", action="Edit", book=form_data, categories=categories, book_id=book_id)

    return render_template("book_form.html", action="Edit", book=book, categories=categories, book_id=book_id)


@app.route("/books/<int:book_id>/delete", methods=["POST"])
def delete_book(book_id):
    success, msg = BookService.delete_book(book_id)
    if success:
        flash(msg, "success")
    else:
        flash(msg, "danger")
    return redirect(url_for("books_list"))


# =============================================================================
# 3. MEMBERS MANAGEMENT
# =============================================================================
@app.route("/members")
def members_list():
    search_query = request.args.get("q", "").strip()
    sort_by = request.args.get("sort", "name")

    members = MemberService.get_all_members(search_query=search_query, sort_by=sort_by)
    return render_template(
        "members.html",
        members=members,
        search_query=search_query,
        selected_sort=sort_by,
        total_members=len(members)
    )


@app.route("/members/new", methods=["GET", "POST"])
def add_member():
    if request.method == "POST":
        member_id = request.form.get("member_id", "").strip()
        name = request.form.get("name", "").strip()
        phone = request.form.get("phone", "").strip()
        email = request.form.get("email", "").strip()

        success, msg, m_id = MemberService.add_member(
            member_id=member_id,
            name=name,
            phone=phone,
            email=email
        )

        if success:
            flash(msg, "success")
            return redirect(url_for("members_list"))
        else:
            flash(msg, "danger")
            return render_template("member_form.html", action="Add", member=request.form)

    # Suggest next member ID
    all_m = MemberService.get_all_members()
    next_id = f"MEM-{101 + len(all_m)}"
    return render_template("member_form.html", action="Add", member={"member_id": next_id})


@app.route("/members/<int:id>/edit", methods=["GET", "POST"])
def edit_member(id):
    member = MemberService.get_member_by_id(id)
    if not member:
        flash("Member not found.", "danger")
        return redirect(url_for("members_list"))

    if request.method == "POST":
        member_id = request.form.get("member_id", "").strip()
        name = request.form.get("name", "").strip()
        phone = request.form.get("phone", "").strip()
        email = request.form.get("email", "").strip()

        success, msg = MemberService.update_member(
            id=id,
            member_id=member_id,
            name=name,
            phone=phone,
            email=email
        )

        if success:
            flash(msg, "success")
            return redirect(url_for("members_list"))
        else:
            flash(msg, "danger")
            form_data = dict(request.form)
            form_data["id"] = id
            return render_template("member_form.html", action="Edit", member=form_data, id=id)

    return render_template("member_form.html", action="Edit", member=member, id=id)


@app.route("/members/<int:id>/delete", methods=["POST"])
def delete_member(id):
    success, msg = MemberService.delete_member(id)
    if success:
        flash(msg, "success")
    else:
        flash(msg, "danger")
    return redirect(url_for("members_list"))


@app.route("/members/<int:id>")
def member_details(id):
    member = MemberService.get_member_by_id(id)
    if not member:
        flash("Member not found.", "danger")
        return redirect(url_for("members_list"))

    loans = MemberService.get_member_loans(id)
    # Enrich active loans with real-time fine calculation
    for loan in loans:
        if loan["status"] == "Issued":
            od, fine = TransactionService.calculate_fine(loan["due_date"])
            loan["overdue_days"] = od
            loan["current_fine"] = fine
            loan["is_overdue"] = od > 0
        else:
            loan["overdue_days"] = 0
            loan["current_fine"] = loan["fine"] or 0.0
            loan["is_overdue"] = False

    return render_template("member_details.html", member=member, loans=loans)


# =============================================================================
# 4. ISSUE BOOK
# =============================================================================
@app.route("/issue", methods=["GET", "POST"])
def issue_book():
    if request.method == "POST":
        book_id_str = request.form.get("book_id", "")
        member_id_str = request.form.get("member_id", "")
        issue_date = request.form.get("issue_date", "").strip()
        due_date = request.form.get("due_date", "").strip()

        try:
            book_id = int(book_id_str)
            member_id = int(member_id_str)
        except ValueError:
            flash("Please select both a valid Book and Member.", "danger")
            return redirect(url_for("issue_book"))

        success, msg, tx_id = TransactionService.issue_book(
            book_id=book_id,
            member_id=member_id,
            issue_date_str=issue_date,
            due_date_str=due_date
        )

        if success:
            flash(msg, "success")
            return redirect(url_for("return_book"))
        else:
            flash(msg, "danger")
            return redirect(url_for("issue_book"))

    available_books = BookService.get_available_books()
    members = MemberService.get_all_members()
    today_str = TransactionService.get_today_str()
    due_date_str = TransactionService.get_default_due_date_str()

    prefill_book_id = request.args.get("book_id", type=int)
    prefill_member_id = request.args.get("member_id", type=int)

    return render_template(
        "issue_book.html",
        available_books=available_books,
        members=members,
        today_str=today_str,
        due_date_str=due_date_str,
        prefill_book_id=prefill_book_id,
        prefill_member_id=prefill_member_id
    )


# =============================================================================
# 5. RETURN BOOK & ACTIVE LOANS
# =============================================================================
@app.route("/return")
def return_book():
    search_query = request.args.get("q", "").strip()
    active_loans = TransactionService.get_active_issued_transactions(search_query=search_query)
    today_str = TransactionService.get_today_str()
    return render_template(
        "return_book.html",
        active_loans=active_loans,
        search_query=search_query,
        today_str=today_str
    )


@app.route("/return/<int:tx_id>", methods=["POST"])
def process_return(tx_id):
    return_date = request.form.get("return_date", "").strip() or TransactionService.get_today_str()
    success, msg, fine, overdue_days = TransactionService.return_book(
        transaction_id=tx_id,
        return_date_str=return_date
    )

    if success:
        flash(msg, "success" if fine == 0 else "warning")
    else:
        flash(msg, "danger")

    # Redirect back to referring page or return list
    next_url = request.form.get("next") or request.referrer or url_for("return_book")
    return redirect(next_url)


# =============================================================================
# 6. TRANSACTIONS MANAGEMENT
# =============================================================================
@app.route("/transactions")
def transactions_list():
    status_filter = request.args.get("status", "All")
    search_query = request.args.get("q", "").strip()

    transactions = TransactionService.get_all_transactions(
        search_query=search_query,
        status=status_filter
    )

    today_str = TransactionService.get_today_str()
    return render_template(
        "transactions.html",
        transactions=transactions,
        status_filter=status_filter,
        search_query=search_query,
        today_str=today_str
    )


@app.route("/transactions/export")
def export_transactions():
    """Export any supported report as CSV, Excel, or PDF."""
    report_type = request.args.get("type", "All Books")
    export_format = request.args.get("format", "csv").lower()
    headers, rows = TransactionService.get_reports_data(report_type)
    safe_name = "_".join(report_type.lower().split())
    stamp = datetime.now().strftime('%Y%m%d_%H%M%S')

    if export_format in ("xlsx", "excel"):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter
        except ImportError:
            flash("Excel export requires openpyxl. Install dependencies from requirements.txt.", "error")
            return redirect(url_for("settings"))
        output = io.BytesIO()
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = report_type[:31] or "Report"
        sheet.append(headers)
        for row in rows:
            sheet.append([str(value) if value is not None else "" for value in row])
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1D4ED8")
            cell.alignment = Alignment(wrap_text=True)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column_cells in sheet.columns:
            max_len = min(max(len(str(c.value or "")) for c in column_cells) + 2, 42)
            sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = max(max_len, 12)
        workbook.save(output)
        output.seek(0)
        return send_file(output, as_attachment=True, download_name=f"{safe_name}_{stamp}.xlsx",
                         mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

    if export_format == "pdf":
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import landscape, A4
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.lib.units import mm
            from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        except ImportError:
            flash("PDF export requires reportlab. Install dependencies from requirements.txt.", "error")
            return redirect(url_for("settings"))
        output = io.BytesIO()
        doc = SimpleDocTemplate(output, pagesize=landscape(A4), rightMargin=12*mm, leftMargin=12*mm,
                                topMargin=12*mm, bottomMargin=12*mm)
        styles = getSampleStyleSheet()
        story = [Paragraph(f"LibraryMS Report: {report_type}", styles["Title"]),
                 Paragraph(f"Generated: {datetime.now().strftime('%d %b %Y, %I:%M %p')}", styles["Normal"]), Spacer(1, 8*mm)]
        table_data = [[str(v) for v in headers]] + [[str(v) if v is not None else "" for v in row] for row in rows]
        if not table_data[0]:
            table_data = [["No report columns available"]]
        available_width = landscape(A4)[0] - 24*mm
        col_width = available_width / max(len(table_data[0]), 1)
        column_widths = [col_width for _ in table_data[0]]
        table = Table(table_data, repeatRows=1, colWidths=column_widths)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1D4ED8")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
            ("LEADING", (0, 0), (-1, -1), 9),
            ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#CBD5E1")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EFF6FF")]),
            ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ]))
        story.append(table)
        doc.build(story)
        output.seek(0)
        return send_file(output, as_attachment=True, download_name=f"{safe_name}_{stamp}.pdf", mimetype="application/pdf")

    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(headers)
    writer.writerows(rows)
    return Response("\ufeff" + output.getvalue(), mimetype="text/csv; charset=utf-8",
                    headers={"Content-Disposition": f"attachment; filename={safe_name}_{stamp}.csv"})


@app.route("/settings/backup")
def download_database_backup():
    """Create a consistent SQLite backup without exposing the live database file directly."""
    if session.get("role", "Admin").lower() != "admin":
        flash("Only an administrator can back up the database.", "error")
        return redirect(url_for("settings"))
    if not os.path.isfile(DB_PATH):
        flash("Database file was not found.", "error")
        return redirect(url_for("settings"))
    temp = tempfile.NamedTemporaryFile(prefix="libraryms_backup_", suffix=".db", delete=False)
    temp_path = temp.name
    temp.close()
    try:
        with sqlite3.connect(DB_PATH) as source, sqlite3.connect(temp_path) as destination:
            source.backup(destination)
        response = send_file(temp_path, as_attachment=True,
                             download_name=f"libraryms_backup_{datetime.now().strftime('%Y%m%d_%H%M%S')}.db",
                             mimetype="application/vnd.sqlite3")
        response.call_on_close(lambda: os.path.exists(temp_path) and os.unlink(temp_path))
        return response
    except Exception:
        try:
            if os.path.exists(temp_path):
                os.unlink(temp_path)
        except OSError:
            pass
        raise


@app.route("/settings/restore", methods=["POST"])
def restore_database_backup():
    """Restore only a valid SQLite database containing the expected LMS tables."""
    if session.get("role", "Admin").lower() != "admin":
        flash("Only an administrator can restore the database.", "error")
        return redirect(url_for("settings"))
    uploaded = request.files.get("backup_file")
    if not uploaded or not uploaded.filename:
        flash("Choose a database backup file first.", "error")
        return redirect(url_for("settings"))
    os.makedirs(os.path.dirname(os.path.abspath(DB_PATH)), exist_ok=True)
    fd, candidate_path = tempfile.mkstemp(prefix="libraryms_restore_", suffix=".db")
    os.close(fd)
    try:
        uploaded.save(candidate_path)
        with sqlite3.connect(candidate_path) as candidate:
            integrity = candidate.execute("PRAGMA integrity_check").fetchone()
            tables = {r[0] for r in candidate.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        required = {"users", "books", "members", "transactions", "settings"}
        if not integrity or integrity[0] != "ok" or not required.issubset(tables):
            flash("Restore cancelled: file is not a valid LibraryMS database backup.", "error")
            return redirect(url_for("settings"))
        # Preserve the current database before replacing it.
        if os.path.exists(DB_PATH):
            pre_restore = DB_PATH + ".pre_restore_backup"
            shutil.copy2(DB_PATH, pre_restore)
        os.replace(candidate_path, DB_PATH)
        flash("Database restored successfully. Your previous database was preserved as .pre_restore_backup.", "success")
    except (OSError, sqlite3.Error) as exc:
        flash(f"Could not restore database: {exc}", "error")
    finally:
        try:
            if os.path.exists(candidate_path):
                os.unlink(candidate_path)
        except OSError:
            pass
    return redirect(url_for("settings"))


# =============================================================================
# 7. SETTINGS & PREFERENCES
# =============================================================================
@app.route("/settings", methods=["GET", "POST"])
def settings():
    if request.method == "POST":
        dark_mode = request.form.get("dark_mode") == "1"
        high_contrast = request.form.get("high_contrast") == "1"
        try:
            font_size = int(request.form.get("font_size", 14))
        except ValueError:
            font_size = 14

        SettingsService.save_settings(dark_mode, high_contrast, font_size)
        session["theme"] = "dark" if dark_mode else "light"
        flash("Settings saved successfully!", "success")
        return redirect(url_for("settings"))

    current_settings = SettingsService.load_settings()
    stats = TransactionService.get_dashboard_stats()

    return render_template(
        "settings.html",
        settings=current_settings,
        stats=stats,
        db_path=DB_PATH
    )


@app.route("/toggle-theme", methods=["POST"])
def toggle_theme():
    """Quick theme switcher called via navbar toggle."""
    data = request.get_json(silent=True) or {}
    theme = data.get("theme")

    if not theme:
        # Toggle current session theme
        current = session.get("theme", "dark")
        theme = "light" if current == "dark" else "dark"

    session["theme"] = theme
    is_dark = (theme == "dark")
    SettingsService.save_settings(is_dark, False, 14)
    return jsonify({"success": True, "theme": theme})


# =============================================================================
# ERROR HANDLERS
# =============================================================================
@app.errorhandler(404)
def not_found_error(error):
    return render_template("base.html", error_title="404 - Page Not Found", error_message="The requested page does not exist."), 404


@app.errorhandler(500)
def internal_error(error):
    return render_template("base.html", error_title="500 - Internal Server Error", error_message="An unexpected server error occurred."), 500


# =============================================================================
# APPLICATION ENTRY POINT
# =============================================================================
# Gunicorn imports this module without executing the __main__ block, so initialize
# the schema here as well. init_db() is designed to be safe to rerun.
init_db()

if __name__ == "__main__":

    host = "127.0.0.1"
    port = 5000

    print("=" * 70)
    print(f"  {APP_TITLE.upper()}")
    print(f"  {APP_SUBTITLE}")
    print("=" * 70)
    print(f"  * Status: Running locally")
    print(f"  * Local Web URL: http://{host}:{port}")
    print(f"  * Open in Google Chrome: http://{host}:{port}")
    print(f"  * Database: {DB_PATH}")
    print("=" * 70)
    print("  Press Ctrl+C to stop the server.\n")

    app.run(host=host, port=port, debug=True)
